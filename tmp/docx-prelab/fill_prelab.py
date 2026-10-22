from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_run_font(run, name="Cambria", size=11, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_paragraph(paragraph, text, *, bold=None):
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, bold=bold)
    return paragraph


def add_identity_before(title_paragraph, label, value):
    paragraph = title_paragraph.insert_paragraph_before()
    paragraph.style = "normal"
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=11, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=11, bold=False)
    return paragraph


def set_cell_text(cell, text):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            set_run_font(run, size=10.5)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def populate_table(table, rows):
    while len(table.rows) < len(rows) + 1:
        table.add_row()
    while len(table.rows) > len(rows) + 1:
        table._tbl.remove(table.rows[-1]._tr)

    set_repeat_table_header(table.rows[0])
    for row in table.rows:
        row.height = None
        row.height_rule = WD_ROW_HEIGHT_RULE.AUTO

    for row_index, values in enumerate(rows, start=1):
        for column_index, value in enumerate(values):
            set_cell_text(table.cell(row_index, column_index), value)


def remove_paragraph(paragraph):
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)


def remove_table(table):
    table._element.getparent().remove(table._element)


def main(source_path: str, output_path: str):
    doc = Document(source_path)

    # Identity block above the template title.
    title = doc.paragraphs[0]
    add_identity_before(title, "Nama", "Rima Marsya Putri Sabrina")
    add_identity_before(title, "NIM", "202310370311425")

    # Requirements: keep the existing heading and body-paragraph pattern.
    requirements_heading = next(p for p in doc.paragraphs if p.text == "1. Requirements")
    idx = next(i for i, p in enumerate(doc.paragraphs) if p.text == "1. Requirements")
    requirement_paragraphs = doc.paragraphs[idx + 1 : idx + 3]
    requirement_texts = [
        "RQ-001: Sistem harus memungkinkan mahasiswa mendaftar mata kuliah yang tersedia.",
        "RQ-002: Sistem harus menolak pendaftaran mata kuliah yang telah mencapai kapasitas maksimum.",
        "RQ-003: Sistem harus mencegah mahasiswa mendaftar mata kuliah yang sama lebih dari satu kali.",
        "RQ-004: Sistem harus memungkinkan mahasiswa melihat daftar mata kuliah yang berhasil didaftarkan.",
        "RQ-005: Sistem harus menampilkan pesan yang sesuai apabila belum ada data registrasi mata kuliah.",
    ]
    replace_paragraph(requirement_paragraphs[0], requirement_texts[0])
    replace_paragraph(requirement_paragraphs[1], requirement_texts[1])
    feature_heading = next(p for p in doc.paragraphs if p.text == "2. Feature Sets")
    for text in requirement_texts[2:]:
        paragraph = feature_heading.insert_paragraph_before(text, style="normal")
        replace_paragraph(paragraph, text)

    # Feature sets: retain template labels and paragraph sequence; answers are Indonesian.
    feature_heading = next(p for p in doc.paragraphs if p.text == "2. Feature Sets")
    feature_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == "2. Feature Sets")
    blocks = doc.paragraphs[feature_index + 1 : feature_index + 9]
    feature_texts = [
        ("FS-001 - Registrasi Mata Kuliah", True),
        ("Objective: Mengelola proses pendaftaran mata kuliah yang tersedia dengan validasi kapasitas dan duplikasi.", False),
        ("Priority: Tinggi", False),
        ("Traceability: RQ-001, RQ-002, RQ-003", False),
        ("FS-002 - Daftar Mata Kuliah Terdaftar", True),
        ("Objective: Menampilkan daftar mata kuliah yang berhasil didaftarkan atau pesan ketika data registrasi belum tersedia.", False),
        ("Priority: Tinggi", False),
        ("Traceability: RQ-004, RQ-005", False),
    ]
    for paragraph, (text, bold) in zip(blocks, feature_texts):
        replace_paragraph(paragraph, text, bold=bold)

    test_conditions = [
        ["TCond-001", "FS-001", "Pendaftaran mata kuliah berhasil ketika mata kuliah tersedia dan kapasitas belum penuh", "Tinggi", "RQ-001"],
        ["TCond-002", "FS-001", "Pendaftaran ditolak ketika mata kuliah telah mencapai kapasitas maksimum", "Tinggi", "RQ-002"],
        ["TCond-003", "FS-001", "Pendaftaran ditolak ketika mahasiswa memilih mata kuliah yang sudah pernah didaftarkan", "Tinggi", "RQ-003"],
        ["TCond-004", "FS-002", "Sistem menampilkan mata kuliah yang berhasil didaftarkan", "Tinggi", "RQ-004"],
        ["TCond-005", "FS-002", "Sistem menampilkan seluruh mata kuliah ketika mahasiswa memiliki lebih dari satu registrasi berhasil", "Sedang", "RQ-004"],
        ["TCond-006", "FS-002", "Sistem menampilkan pesan bahwa data registrasi tidak tersedia ketika belum ada mata kuliah yang didaftarkan", "Tinggi", "RQ-005"],
    ]
    coverage_items = [
        ["CI-001", "Ketersediaan mata kuliah dan keberhasilan registrasi", "TCond-001", "Tinggi", "RQ-001"],
        ["CI-002", "Status kapasitas maksimum mata kuliah", "TCond-002", "Tinggi", "RQ-002"],
        ["CI-003", "Status duplikasi pendaftaran mata kuliah", "TCond-003", "Tinggi", "RQ-003"],
        ["CI-004", "Data mata kuliah yang berhasil didaftarkan", "TCond-004", "Tinggi", "RQ-004"],
        ["CI-005", "Kelengkapan daftar mata kuliah terdaftar", "TCond-005", "Sedang", "RQ-004"],
        ["CI-006", "Kondisi daftar kosong dan pesan tidak ada data registrasi", "TCond-006", "Tinggi", "RQ-005"],
    ]
    ttm_rows = [
        ["RQ-001", "FS-001", "TCond-001", "CI-001"],
        ["RQ-002", "FS-001", "TCond-002", "CI-002"],
        ["RQ-003", "FS-001", "TCond-003", "CI-003"],
        ["RQ-004", "FS-002", "TCond-004", "CI-004"],
        ["RQ-004", "FS-002", "TCond-005", "CI-005"],
        ["RQ-005", "FS-002", "TCond-006", "CI-006"],
    ]
    populate_table(doc.tables[0], test_conditions)
    populate_table(doc.tables[1], coverage_items)
    populate_table(doc.tables[2], ttm_rows)

    # Keep the first large table with its heading instead of leaving an orphaned
    # heading at the bottom of page 1.
    test_conditions_heading = next(p for p in doc.paragraphs if p.text == "3. Test Conditions")
    test_conditions_heading.paragraph_format.page_break_before = True

    # Remove the sample Section 6 because this module's deliverable ends at Section 5.
    section_six = next((p for p in doc.paragraphs if p.text == "6. Test Design Technique Summary"), None)
    if len(doc.tables) > 3:
        remove_table(doc.tables[3])
    if section_six is not None:
        remove_paragraph(section_six)

    # Remove trailing empty paragraphs left by the sample section.
    for paragraph in list(doc.paragraphs):
        if not paragraph.text.strip() and paragraph._element.getnext() is None:
            remove_paragraph(paragraph)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Usage: fill_prelab.py SOURCE_DOCX OUTPUT_DOCX")
    main(sys.argv[1], sys.argv[2])
